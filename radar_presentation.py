"""Editorial briefing: weekly signals, IOMA, Google queries, clickable sources."""
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from campaign import logo_layer,wrap
from studio import font

INK=(38,35,32);MUTED=(100,94,87);PAPER=(249,247,242);RULE=(195,186,173)


def text(slide,value,x,y,w,h,size=16,serif=False,color=INK,link=None):
    f=font(size,serif)
    lines=[line for part in value.split('\n') for line in (wrap(part,f,w*72) or [''])]
    if len(lines)*size*1.23>h*72 or any(f.getlength(line)>w*72 for line in lines):
        raise ValueError('Testo troppo lungo per la slide: '+value[:65])
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    frame=shape.text_frame
    frame.margin_top=frame.margin_bottom=frame.margin_left=frame.margin_right=0
    frame.word_wrap=True;frame.auto_size=MSO_AUTO_SIZE.NONE
    frame.text=value
    for p in frame.paragraphs:
        p.font.name='Cormorant Garamond' if serif else 'DejaVu Sans'
        p.font.size=Pt(size);p.font.color.rgb=RGBColor(*color)
        p.line_spacing=1.08;p.space_before=p.space_after=Pt(0)
    if link:
        shape.click_action.hyperlink.address=link
    return shape


def rule(slide,x,y,w):
    s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(.012))
    s.fill.solid();s.fill.fore_color.rgb=RGBColor(*RULE);s.line.fill.background()
    s._element.spPr.append(OxmlElement('a:effectLst'))
    for effect in s._element.xpath('./p:style/a:effectRef'):
        effect.set('idx','0')


def image(slide,path,x,y,w,h):
    with Image.open(path) as im: ratio=min(w/im.width,h/im.height);pw,ph=im.width*ratio,im.height*ratio
    slide.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))


def page(deck,number,lang,period):
    slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(*PAPER)
    text(slide,'BEAUTY RADAR',.6,.42,4,.3,size=11)
    stream=BytesIO();logo_layer(220,False).save(stream,format='PNG');stream.seek(0)
    slide.shapes.add_picture(stream,Inches(6.1167),Inches(.24),width=Inches(1.1))
    text(slide,period,9.25,.42,3.45,.3,size=11,color=MUTED)
    rule(slide,.6,.92,12.1)
    text(slide,'IOMA PARIS  /  '+lang.upper(),.6,7.14,9,.22,size=9,color=MUTED)
    text(slide,f'{number:02}',12.12,7.12,.5,.24,size=10,color=MUTED)
    return slide


def article_sources(deck,brief,lang,period):
    fr=lang=='fr'
    slide=page(deck,4,lang,period)
    text(slide,'Articles & sources' if fr else 'Articoli e fonti',.6,1.15,12.1,.65,size=34,serif=True)
    text(slide,'Ouvrir les publications citées dans ce radar' if fr else 'Apri le pubblicazioni citate nel radar',.6,1.87,12,.4,size=15,color=MUTED)
    names={'editorial':('ARTICLE ÉDITORIAL','ARTICOLO EDITORIALE'),
           'advertising_campaign':('PRESSE MARKETING · CAMPAGNE','STAMPA MARKETING · CAMPAGNA'),
           'brand_promotion':('COMMUNIQUÉ DE MARQUE','COMUNICATO DEL MARCHIO')}
    order={'editorial':0,'advertising_campaign':1,'brand_promotion':2}
    for index,source in enumerate(sorted(brief['sources'],key=lambda s:order[s['kind']])):
        y=2.65+index*1.24
        date=datetime.fromisoformat(source['published_at']).strftime('%d.%m.%Y')
        text(slide,source['publisher']+'  ·  '+date+'  ·  '+names[source['kind']][0 if fr else 1],.6,y,12.1,.3,size=11,color=MUTED)
        title=source.get('display_title',{}).get(lang,source['title'])
        text(slide,title,.6,y+.34,9.5,.58,size=23,serif=True)
        label=('Lire le communiqué' if source['kind']=='brand_promotion' else 'Lire l’article') if fr else ('Leggi il comunicato' if source['kind']=='brand_promotion' else 'Leggi l’articolo')
        shape=text(slide,label,10.2,y+.39,2.5,.4,size=12,link=source['url'])
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                run.hyperlink.address=source['url'];run.font.underline=True
        rule(slide,.6,y+1.04,12.1)
    text(slide,'Date de publication indiquée ; période des campagnes précisée en page 1.' if fr else 'Data di pubblicazione indicata; periodo delle campagne precisato a pagina 1.',.6,6.6,12,.35,size=11,color=MUTED)
    slide.notes_slide.notes_text_frame.text=json.dumps(brief['sources'],ensure_ascii=False,indent=2)


def google_queries(deck,brief,lang):
    fr=lang=='fr';result=brief['google_trends_results'][lang]
    data=result['data']
    period=' — '.join(datetime.fromisoformat(data[key]).strftime('%d.%m.%Y') for key in ('period_start','period_end'))
    slide=page(deck,3,lang,period)
    text(slide,'Les recherches beauté sur Google' if fr else 'Le ricerche beauty su Google',.6,1.15,12.1,.65,size=32,serif=True)
    text(slide,'France · Soins du visage et du corps · Recherche Web' if fr else 'Italia · Cura di viso e corpo · Ricerca Web',.6,1.87,12,.4,size=15,color=MUTED)
    text(slide,'REQUÊTES PRINCIPALES · INDICE 0–100' if fr else 'QUERY PRINCIPALI · INDICE 0–100',.6,2.55,6.2,.33,size=12)
    text(slide,'EN CROISSANCE · SÉLECTION BEAUTÉ' if fr else 'IN CRESCITA · SELEZIONE BEAUTY',7.2,2.55,5.5,.33,size=12)
    rule(slide,.6,2.99,5.7);rule(slide,7.2,2.99,5.5)
    for index,row in enumerate(data['top']):
        y=3.17+index*.51
        text(slide,row['query'],.6,y,4.7,.39,size=19)
        text(slide,str(row['value']),5.45,y,.85,.37,size=17)
    for index,row in enumerate(data['rising']):
        y=3.17+index*.8
        text(slide,row['query'],7.2,y,3.8,.68,size=15)
        value=('Percée' if fr else 'Impennata') if row['breakout'] else '+'+str(row['value'])+' %'
        text(slide,value,11.15,y,1.55,.45,size=14)
    comparison='–'.join(datetime.fromisoformat(data[key]).strftime('%d.%m.%Y') for key in ('comparison_start','comparison_end'))
    text(slide,('Indice relatif, pas un volume. Croissance vs ' if fr else 'Indice relativo, non volume. Crescita rispetto al ')+comparison+'.',.6,5.88,12.1,.32,size=11,color=MUTED)
    text(slide,data['interpretation'],.6,6.3,12.1,.48,size=13)
    text(slide,'Google Trends · ouvrir les données' if fr else 'Google Trends · apri i dati',.6,6.84,4.5,.23,size=10,color=MUTED,link=data['source_url'])
    text(slide,'Sélection pertinente ; ordre Google conservé. Méthode ↗' if fr else 'Selezione pertinente; ordine Google conservato. Metodo ↗',5.4,6.84,7.3,.23,size=10,color=MUTED,link=data['method_url'])
    slide.notes_slide.notes_text_frame.text=json.dumps(result,ensure_ascii=False,indent=2)


def create_presentation(product,report,graphics,destination,lang):
    brief=report['weekly'];ingredients=report['ingredients'];copy=report['localizations'][lang]
    fr=lang=='fr'
    end=datetime.fromisoformat(report['created_at']).strftime('%d.%m.%Y')
    start=datetime.fromisoformat(report['window_start']).strftime('%d.%m.%Y')
    deck=Presentation();deck.slide_width=Inches(13.333333);deck.slide_height=Inches(7.5)
    deck.core_properties.title='IOMA Beauty Radar — '+lang.upper()
    deck.core_properties.author='IOMA Paris'
    slide=page(deck,1,lang,start+' — '+end)
    text(slide,'La semaine skincare' if fr else 'La settimana skincare',.6,1.15,12.1,.65,size=34,serif=True)
    text(slide,'Signaux éditoriaux, actifs et campagnes observées' if fr else 'Segnali editoriali, attivi e campagne osservate',.6,1.87,12,.4,size=15,color=MUTED)
    columns=[.6,4.73,8.86]
    headings=['01  SKINCARE','02  INGRÉDIENTS' if fr else '02  INGREDIENTI','03  PUBLICITÉ & PROMOTIONS' if fr else '03  PUBBLICITÀ & PROMOZIONI']
    for x,title in zip(columns,headings):
        text(slide,title,x,2.62,3.83,.38,size=12);rule(slide,x,3.13,3.83)
    labels=' + '.join(i['label'][lang].replace('\n',' ') for i in ingredients)
    text(slide,'Barrière & fermeté' if fr else 'Barriera & compattezza',.6,3.36,3.83,.75,size=26,serif=True)
    text(slide,brief['insight'][lang],.6,4.24,3.83,1.18,size=17)
    editorial=next((s for s in brief['sources'] if s['kind']=='editorial'),None)
    if editorial:
        text(slide,editorial['summary'][lang],.6,5.62,3.83,.55,size=14)
    text(slide,labels,4.73,3.36,3.83,1.1,size=25,serif=True)
    rationale=('Deux familles présentes dans la formule IOMA et citées dans les sources récentes. Sélection éditoriale, pas un classement par concentration.' if fr else
               'Due famiglie presenti nella formula IOMA e citate nelle fonti recenti. Selezione editoriale, non una classifica per concentrazione.')
    text(slide,rationale,4.73,4.68,3.83,1.38,size=16)
    campaigns=[s for s in brief['sources'] if s['kind']!='editorial']
    for i,s in enumerate(campaigns[:2]):
        text(slide,s['summary'][lang],8.86,3.38+i*1.25,3.83,1.17,size=14)
    text(slide,'Lecture internationale ; diffusion payante et budgets non audités.' if fr else 'Osservazione internazionale; erogazione paid e budget non verificati.',8.86,6.08,3.83,.6,size=11,color=MUTED)
    for i,s in enumerate(brief['sources']):
        date=datetime.fromisoformat(s['published_at']).strftime('%d/%m')
        text(slide,f'[{i+1}] {s["publisher"]} · {date}',.6+i*4.13,6.79,3.83,.23,size=9,color=MUTED,link=s['url'])
    slide.notes_slide.notes_text_frame.text=json.dumps(brief,ensure_ascii=False,indent=2)

    slide=page(deck,2,lang,start+' — '+end)
    text(slide,copy['product_name'],.6,1.13,12.1,.62,size=32,serif=True)
    text(slide,'Des tendances aux réseaux sociaux' if fr else 'Dai trend ai contenuti social',.6,1.87,12,.4,size=15,color=MUTED)
    image(slide,graphics['feed'],.6,2.53,3.2,4)
    strategy=brief['strategy'][lang]
    text(slide,'LE PRODUIT & LE LIEN' if fr else 'IL PRODOTTO E IL COLLEGAMENTO',4.22,2.59,7.9,.33,size=12)
    text(slide,strategy['positioning'],4.22,3.07,8.02,.98,size=17)
    text(slide,strategy['audience'],4.22,4.12,8.02,.6,size=14,color=MUTED)
    rule(slide,4.22,4.84,8.45)
    text(slide,'ACTIVATION À TESTER' if fr else 'ATTIVAZIONE DA TESTARE',4.22,5.02,8.02,.3,size=12)
    text(slide,strategy['activation'],4.22,5.45,8.02,.77,size=14)
    text(slide,strategy['measurement'],4.22,6.32,8.02,.63,size=12,color=MUTED)
    text(slide,'Formule & INCI IOMA' if fr else 'Formula e INCI IOMA',.6,6.76,3.2,.24,size=10,color=MUTED,link=report['creative']['source'])
    slide.notes_slide.notes_text_frame.text=json.dumps({'ingredients':ingredients,'caption':copy['caption'],'strategy':strategy,'source':report['creative']['source']},ensure_ascii=False,indent=2)
    google_queries(deck,brief,lang)
    article_sources(deck,brief,lang,start+' — '+end)
    deck.save(destination)
    return {'path':str(destination),'slides':4,'language':lang,
            'source_urls':[s['url'] for s in brief['sources']]+[brief['google_trends_results'][lang]['data']['source_url']]}
