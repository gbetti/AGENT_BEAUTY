from argparse import Namespace
from datetime import datetime, timezone
from pathlib import Path
import hashlib
import json
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
from PIL import Image, ImageChops
from agent import load_products, parse_news, run
from campaign import FORMATS, packshot_layer
from editorial import BRIEFS, LANGUAGES, campaign_plan
from studio import render_campaign, validate
from pptx import Presentation
from weekly import load_brief, select_ingredients
from google_trends import select_queries, research_trends, week_window

ROOT=Path(__file__).resolve().parents[1]

def weekly_fixture():
    brief=json.loads((ROOT/'research/2026-10-08.json').read_text())
    for source in brief['sources']: source['in_week']=source['id']!='cerave'
    return brief

def trends_fixture():
    now=datetime(2026,10,8,12,tzinfo=timezone.utc)
    brief=weekly_fixture();results={}
    for lang,geo in [('fr','FR'),('it','IT')]:
        packet=json.loads((ROOT/'research/google-trends/2026-10-01_2026-10-07'/(geo+'.json')).read_text())
        results[lang]={'data':select_queries(packet,brief['google_trends']['markets'][lang],now),'evidence':packet}
    return results

def ingredient_fixture():
    return json.loads((ROOT/'assets/ingredients.json').read_text())['creme-sublime-revitalisante'][:2]


class AgentTests(unittest.TestCase):
    def test_news_filters_old_future_invalid_and_duplicate(self):
        xml=b'''<rss><channel>
        <item><title>Glow skincare</title><link>https://example.com/one</link><pubDate>Wed, 07 Oct 2026 10:00:00 GMT</pubDate></item>
        <item><title>Glow skincare</title><link>https://example.com/one</link><pubDate>Wed, 07 Oct 2026 10:00:00 GMT</pubDate></item>
        <item><title>Old glow</title><link>https://example.com/old</link><pubDate>Wed, 23 Sep 2026 10:00:00 GMT</pubDate></item>
        <item><title>Future glow</title><link>https://example.com/future</link><pubDate>Wed, 14 Oct 2026 10:00:00 GMT</pubDate></item>
        <item><title>George returns home</title><link>https://example.com/off-topic</link><pubDate>Wed, 07 Oct 2026 11:00:00 GMT</pubDate></item>
        </channel></rss>'''
        self.assertEqual(len(parse_news(xml,datetime(2026,10,7,12,tzinfo=timezone.utc))),1)

    def test_original_assets_are_unchanged_and_high_resolution(self):
        for asset in json.loads((ROOT/'assets/products/manifest.json').read_text()):
            path=ROOT/asset['path']
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),asset['sha256'])
            with Image.open(path) as im: self.assertEqual(im.size,(2000,2000))

    def test_every_product_copy_and_scene_fits_all_formats_without_collisions(self):
        scenes=json.loads((ROOT/'assets/backgrounds/editorial-scenes.json').read_text())
        with tempfile.TemporaryDirectory() as folder:
            products=load_products(Path(folder)/'db.sqlite')
            for product in products:
                photo=packshot_layer(product['images'][0]['path'])
                brief=BRIEFS[product['id'].split(':')[-1]]
                plan=campaign_plan(product,[])
                # Every copy in both languages on both light/background variants.
                for variant in brief['headlines']:
                    for language in LANGUAGES:
                        for scene in scenes:
                            with self.subTest(product=product['name'],headline=variant[language],scene=scene['id']):
                                copy={**plan['copies'][language],'headline':variant[language]}
                                if product['id'].endswith('creme-sublime-revitalisante'):
                                    copy['ingredients']=ingredient_fixture()
                                with patch('studio.packshot_layer',return_value=photo):
                                    outputs=render_campaign(product,copy,Path(folder),scene)
                                self.assertEqual(set(outputs),set(FORMATS))
                                for kind,info in outputs.items():
                                    with Image.open(info['path']) as im: self.assertEqual(im.size,FORMATS[kind])
                                    validate(kind,info['elements'])
                                    if copy.get('ingredients'):
                                        self.assertEqual(len([e for e in info['elements'] if e['kind']=='arrow']),2)
                                        self.assertEqual({e['ingredient_id'] for e in info['elements'] if e['kind']=='ingredient'}, {'ceramides','peptides'})
                                    self.assertLess(info['text_area_ratio'],.20)
                                    logo=next(e['box'] for e in info['elements'] if e['kind']=='logo')
                                    self.assertAlmostEqual((logo[0]+logo[2])/2,FORMATS[kind][0]/2)
                                    for e in info['elements']:
                                        if 'text' in e: self.assertGreaterEqual(e['contrast_ratio'],4.5)
                                    if kind!='banner': self.assertNotIn('cta',[e['kind'] for e in info['elements']])

    def test_packshot_label_pixels_are_preserved_at_final_size(self):
        with tempfile.TemporaryDirectory() as folder:
            product=load_products(Path(folder)/'db.sqlite')[-1]
            plan=campaign_plan(product,[])
            plan['copies']['fr']['ingredients']=ingredient_fixture()
            photo=packshot_layer(product['images'][0]['path'])
            with patch('studio.packshot_layer',return_value=photo):
                outputs=render_campaign(product,plan['copies']['fr'],Path(folder),plan['scene'])
            for info in outputs.values():
                box=next(e['box'] for e in info['elements'] if e['kind']=='product')
                l,t,r,b=box
                expected=photo.resize((r-l,b-t),Image.Resampling.LANCZOS)
                with Image.open(info['path']) as im: actual=im.crop(box)
                opaque=expected.getchannel('A').point(lambda v:255 if v==255 else 0)
                delta=ImageChops.difference(actual,expected.convert('RGB'))
                delta.paste((0,0,0),mask=ImageChops.invert(opaque))
                self.assertIsNone(delta.getbbox())

    def test_changed_claim_and_missing_brief_fail_closed(self):
        with tempfile.TemporaryDirectory() as folder:
            product=load_products(Path(folder)/'db.sqlite')[0]
            product['data']['editorial']['claim_it']='Unverified new claim'
            with self.assertRaisesRegex(ValueError,'claim'): campaign_plan(product,[])
            product['id']='unknown'
            with self.assertRaisesRegex(ValueError,'brief'): campaign_plan(product,[])

    def test_consecutive_campaigns_change_scene_and_copy(self):
        with tempfile.TemporaryDirectory() as folder:
            product=load_products(Path(folder)/'db.sqlite')[-1]
            first=campaign_plan(product,[])
            second=campaign_plan(product,[{'product_id':product['id'],'creative':first}])
            self.assertNotEqual(first['variant'],second['variant'])
            self.assertNotEqual(first['scene']['id'],second['scene']['id'])

    def test_six_images_two_decks_and_fresh_research_each_run(self):
        with tempfile.TemporaryDirectory() as folder:
            args=Namespace(database=Path(folder)/'db.sqlite',output=Path(folder)/'out',product='Sublime')
            news=[{'title':'Skin barrier skincare trend','published_at':datetime.now(timezone.utc).isoformat(),'url':'https://example.com/skin'}]
            with patch('agent.search_news',return_value=news) as search, patch('agent.load_brief',side_effect=lambda *a:weekly_fixture()) as weekly_search, patch('agent.select_ingredients',side_effect=lambda *a:ingredient_fixture()) as formula, patch('agent.research_trends',side_effect=lambda *a:trends_fixture()) as trends:
                first=run(args);second=run(args)
            self.assertEqual(search.call_count,2)
            self.assertEqual(weekly_search.call_count,2)
            self.assertEqual(formula.call_count,2)
            self.assertEqual(trends.call_count,2)
            self.assertNotEqual(first,second)
            for d in args.output.iterdir():
                self.assertEqual({p.name for p in d.iterdir()},{'fr','it'})
                self.assertEqual({str(p.relative_to(d)) for p in d.rglob('*') if p.is_file()},
                                 {lang+'/'+kind+'.png' for lang in LANGUAGES for kind in FORMATS} | {lang+'/presentazione.pptx' for lang in LANGUAGES})
                for lang in LANGUAGES:
                    deck=Presentation(d/lang/'presentazione.pptx')
                    self.assertEqual(len(deck.slides),4)
                    links={shape.click_action.hyperlink.address for slide in deck.slides for shape in slide.shapes if shape.has_text_frame and shape.click_action.hyperlink.address}
                    self.assertTrue({s['url'] for s in weekly_fixture()['sources']}<=links)
                    google=trends_fixture()[lang]['data']
                    self.assertIn(google['source_url'],links)
                    page_text=' '.join(sh.text for sh in deck.slides[2].shapes if sh.has_text_frame)
                    self.assertIn('France' if lang=='fr' else 'Italia',page_text)
                    for row in google['top']+google['rising']:self.assertIn(row['query'],page_text)
                    article_links={r.hyperlink.address for sh in deck.slides[3].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs if r.hyperlink.address}
                    self.assertTrue({s['url'] for s in weekly_fixture()['sources']}<=article_links)
            with sqlite3.connect(args.database) as db:
                rows=db.execute('SELECT research_json FROM generations').fetchall()
            self.assertEqual(len(rows),2)
            for raw, in rows:
                report=json.loads(raw)
                self.assertEqual(set(report['localizations']),{'fr','it'})
                self.assertEqual(report['quality']['visual_review'],'pending')
                self.assertEqual(report['quality']['publication_status'],'not_published')

    def test_second_language_failure_rolls_back_all_new_files_and_history(self):
        with tempfile.TemporaryDirectory() as folder:
            args=Namespace(database=Path(folder)/'db.sqlite',output=Path(folder)/'out',product='Sublime')
            args.output.mkdir();prior=args.output/'prior.png';prior.write_bytes(b'keep')
            def partial(product,copy,directory,scene):
                (directory/'feed.png').write_bytes(b'partial')
                if copy['language']=='it': raise ValueError('Italian failure')
                return {'feed':{'path':directory/'feed.png'}}
            with patch('agent.search_news',return_value=[{'title':'Skincare hydration'}]), patch('agent.load_brief',return_value=weekly_fixture()), patch('agent.select_ingredients',return_value=ingredient_fixture()), patch('agent.research_trends',return_value=trends_fixture()), patch('agent.render_campaign',side_effect=partial):
                with self.assertRaisesRegex(ValueError,'Italian failure'): run(args)
            self.assertEqual(list(args.output.iterdir()),[prior])
            with sqlite3.connect(args.database) as db:
                self.assertEqual(db.execute('SELECT COUNT(*) FROM generations').fetchone()[0],0)

    def test_network_failure_exports_nothing(self):
        with tempfile.TemporaryDirectory() as folder:
            args=Namespace(database=Path(folder)/'db.sqlite',output=Path(folder)/'out',product=None)
            with patch('agent.search_news',side_effect=RuntimeError('Network unavailable')):
                with self.assertRaises(RuntimeError): run(args)
            self.assertFalse(args.output.exists())

    def test_second_deck_failure_rolls_back_images_and_first_deck(self):
        with tempfile.TemporaryDirectory() as folder:
            args=Namespace(database=Path(folder)/'db.sqlite',output=Path(folder)/'out',product='Sublime')
            args.output.mkdir();prior=args.output/'prior.png';prior.write_bytes(b'keep')
            def graphics(product,copy,directory,scene):
                result={}
                for kind in FORMATS:
                    path=directory/(kind+'.png');path.write_bytes(b'image fixture')
                    result[kind]={'path':path}
                return result
            def deck(product,report,images,path,lang):
                path.write_bytes(b'deck fixture')
                if lang=='it':raise ValueError('Italian deck failure')
                return {'path':str(path),'slides':2}
            with patch('agent.search_news',return_value=[{'title':'Skincare hydration'}]), patch('agent.load_brief',return_value=weekly_fixture()), patch('agent.select_ingredients',return_value=ingredient_fixture()), patch('agent.research_trends',return_value=trends_fixture()), patch('agent.render_campaign',side_effect=graphics), patch('agent.create_presentation',side_effect=deck):
                with self.assertRaisesRegex(ValueError,'Italian deck failure'):run(args)
            self.assertEqual(list(args.output.iterdir()),[prior])
            with sqlite3.connect(args.database) as db:
                self.assertEqual(db.execute('SELECT COUNT(*) FROM generations').fetchone()[0],0)

    def test_overflow_is_not_silently_reduced(self):
        from studio import draw_text
        with self.assertRaisesRegex(ValueError,'troppo largo'):
            draw_text(Image.new('RGB',(300,250),'white'),'A title that cannot possibly fit',20,70,130,33,'headline',serif=True)


class WeeklyEvidenceTests(unittest.TestCase):
    now=datetime(2026,10,8,10,tzinfo=timezone.utc)

    def load(self,brief):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'brief.json';path.write_text(json.dumps(brief))
            texts={s['url']:' '.join(s['evidence_terms']) for s in brief['sources']}
            with patch('weekly.fetch_text',side_effect=texts.__getitem__):
                return load_brief(path,self.now)

    def test_old_active_campaign_is_distinct_from_current_news(self):
        brief=self.load(weekly_fixture())
        self.assertEqual([s['in_week'] for s in brief['sources']],[True,True,False])
        self.assertTrue(all(s['checked_at'] and len(s['content_sha256'])==64 for s in brief['sources']))

    def test_expired_brief_is_not_reused(self):
        brief=weekly_fixture();brief['reviewed_at']='2026-10-06T00:00:00+00:00'
        with self.assertRaisesRegex(ValueError,'aggiornare'):self.load(brief)

    def test_future_sources_and_inactive_old_campaign_are_rejected(self):
        for mutation in ('future','inactive'):
            brief=weekly_fixture()
            if mutation=='future':brief['sources'][0]['published_at']='2026-10-09T00:00:00+00:00'
            else:brief['sources'][-1]['event_end']='2026-10-07'
            with self.subTest(mutation=mutation),self.assertRaises(ValueError):self.load(brief)

    def test_promotion_is_not_a_substitute_for_advertising(self):
        brief=weekly_fixture();brief['sources']=brief['sources'][:2]
        with self.assertRaisesRegex(ValueError,'pubblicitaria'):self.load(brief)

    def test_only_verified_inci_and_current_topics_select_ingredients(self):
        product={'id':'ioma:creme-sublime-revitalisante'}
        inci='AQUA (WATER), GLYCERIN, CERAMIDE NP, CERAMIDE AP, PALMITOYL TETRAPEPTIDE-7, PARFUM (FRAGRANCE)'
        with patch('weekly.fetch_text',return_value=inci):
            selected=select_ingredients(product,weekly_fixture(),'https://example.com/product',self.now)
        self.assertEqual([i['id'] for i in selected],['ceramides','peptides'])
        self.assertTrue(all(i['source_ids']==['cosrx','byoma'] for i in selected))
        # An ingredient mentioned elsewhere on the page is not formula evidence.
        with patch('weekly.fetch_text',return_value=inci.replace('PALMITOYL TETRAPEPTIDE-7, ','')+' Related products: PALMITOYL TETRAPEPTIDE-7'):
            with self.assertRaisesRegex(ValueError,'due ingredienti'):
                select_ingredients(product,weekly_fixture(),'https://example.com/product',self.now)
        brief=weekly_fixture()
        for source in brief['sources']:source['topics']=[]
        with patch('weekly.fetch_text',return_value=inci):
            with self.assertRaisesRegex(ValueError,'due ingredienti'):
                select_ingredients(product,brief,'https://example.com/product',self.now)


class GoogleTrendsTests(unittest.TestCase):
    now=datetime(2026,10,8,12,tzinfo=timezone.utc)

    def packet(self,geo='FR'):
        return json.loads((ROOT/'research/google-trends/2026-10-01_2026-10-07'/(geo+'.json')).read_text())

    def test_completed_week_excludes_today_and_handles_year_boundary(self):
        self.assertEqual(week_window(self.now),('2026-10-01','2026-10-07'))
        self.assertEqual(week_window(datetime(2027,1,3,tzinfo=timezone.utc)),('2026-12-27','2027-01-02'))

    def test_selection_keeps_google_spelling_scores_and_original_ranks(self):
        result=select_queries(self.packet(),weekly_fixture()['google_trends']['markets']['fr'],self.now)
        self.assertEqual([r['query'] for r in result['top']],['parfum','sephora','ongle','nocibe','yves rocher'])
        self.assertEqual([r['value'] for r in result['top']],[100,43,23,21,15])
        self.assertEqual([r['source_rank'] for r in result['top']],[1,2,3,4,7])
        self.assertEqual([r['value'] for r in result['rising']],[700,350,50])
        self.assertNotIn('parking',str(result))

    def test_wrong_market_week_category_and_seed_are_rejected(self):
        for field,value in [('geo','IT'),('period_start','2026-09-24'),('category_id',44),('seed_keyword','skincare')]:
            packet=self.packet();packet[field]=value
            with self.subTest(field=field),self.assertRaises(ValueError):
                select_queries(packet,weekly_fixture()['google_trends']['markets']['fr'],self.now)
        packet=self.packet();packet['widget_request']['trendinessSettings']['compareTime']='2026-09-01 2026-09-07'
        with self.assertRaises(ValueError):select_queries(packet,weekly_fixture()['google_trends']['markets']['fr'],self.now)

    def test_missing_keyword_expired_capture_and_invalid_score_fail_closed(self):
        for mutation in ('missing','expired','score'):
            packet=self.packet()
            if mutation=='missing':packet['response']['default']['rankedList'][0]['rankedKeyword']=[]
            elif mutation=='expired':packet['retrieved_at']='2026-10-06T00:00:00+00:00'
            else:packet['response']['default']['rankedList'][0]['rankedKeyword'][0]['value']=101
            with self.subTest(mutation=mutation),self.assertRaises(ValueError):
                select_queries(packet,weekly_fixture()['google_trends']['markets']['fr'],self.now)

    def test_each_language_uses_its_own_market(self):
        with patch('google_trends.fetch_market',side_effect=lambda geo,*args:self.packet(geo)) as fetch:
            result=research_trends(weekly_fixture(),self.now)
        self.assertEqual(fetch.call_count,2)
        self.assertEqual(result['fr']['data']['geo'],'FR')
        self.assertEqual(result['it']['data']['geo'],'IT')
        self.assertEqual(result['it']['data']['top'][1]['query'],'crema')
        self.assertEqual(result['it']['data']['top'][1]['value'],89)
