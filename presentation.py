"""Editable PowerPoint campaign with dated evidence and a social post mockup."""
from datetime import datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from campaign import face, logo_layer, wrap
from research import is_beauty_headline

MOCHA = (47, 32, 22)
IVORY = (249, 245, 239)
GOLD = (172, 132, 85)
INK = (38, 29, 24)
MUTED = (107, 95, 83)
WHITE = (255, 255, 255)
TREND_TOPICS = (
    ('K-beauty e rituali coreani', ('k-beauty', 'corean', 'korean')),
    ('Éclat e make-up naturale', ('glow', 'radiance', 'luminos', 'skin tint', 'tinted', 'makeup', 'make-up', 'teint', 'cipria', 'trucco', 'blush')),
    ('Idratazione e barriera cutanea', ('hydrat', 'idrat', 'barrier', 'barriera')),
    ('Longevità e cura anti-età', ('longevity', 'aging', 'ageing', 'collagen', 'collag', 'firm')),
)


def post_copy(product, headline, cta, claim):
    """Use the catalog's existing benefit; do not invent claims from news titles."""
    caption = f"{headline}\n\n{product['name']}\n{claim}\n\n{cta}"
    if not cta.endswith(('.', '!', '?')):
        caption += '.'
    tags = ['#IOMAParis', '#Skincare']
    if product['data'].get('editorial', {}).get('theme') == 'glow':
        tags += ['#CCGel', '#IncarnatoLuminoso', '#MakeupNaturale', '#BeautyRoutine']
    elif 'yeux' in product['name'].casefold():
        tags += ['#ContornoOcchi', '#AntiAge', '#CuraDellaPelle', '#BeautyRoutine']
    else:
        tags += ['#AntiAge', '#CremaViso', '#CuraDellaPelle', '#BeautyRoutine']
    return caption, tags


def weekly_highlights(report):
    """Select up to three different editorial signals with verifiable dates."""
    start = datetime.fromisoformat(report['window_start'])
    end = datetime.fromisoformat(report['created_at'])
    candidates = []
    seen = set()
    for article in report['news']:
        try:
            published = datetime.fromisoformat(article['published_at'])
            valid = published.tzinfo is not None and start <= published <= end
        except (KeyError, TypeError, ValueError):
            continue
        title, url = article.get('title', '').strip(), article.get('url', '')
        if not valid or not title or not is_beauty_headline(title) or not url.startswith('https://') or url in seen:
            continue
        seen.add(url)
        topic = next((label for label, words in TREND_TOPICS
                      if any(word in title.casefold() for word in words)), 'Altri segnali beauty')
        candidates.append({**article, 'topic': topic,
                           'publisher': article.get('publisher') or urlparse(url).netloc})
    candidates.sort(key=lambda article: datetime.fromisoformat(article['published_at']), reverse=True)
    selected = []
    for article in candidates:
        if article['topic'] != 'Altri segnali beauty' and article['topic'] not in {item['topic'] for item in selected}:
            selected.append(article)
        if len(selected) == 3:
            break
    for article in candidates:
        if len(selected) == 3:
            break
        if article not in selected:
            selected.append(article)
    if not selected:
        raise ValueError('Nessuna fonte datata negli ultimi sette giorni per la slide dei trend.')
    return selected


def rectangle(slide, x, y, width, height, fill, border=None, rounded=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill)
    if border:
        shape.line.color.rgb = RGBColor(*border)
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = 0.05
    return shape


def text(slide, value, x, y, width, height, size=20, bold=False, color=INK,
         minimum=11, link=None, center=False):
    # Match the actual font metrics rather than relying on reader-side auto-fit.
    for fitted in range(size, minimum-1, -1):
        font = face(fitted, bold)
        wrapped = [line for part in value.split('\n') for line in (wrap(part, font, width*72) or [''])]
        if len(wrapped)*fitted*1.22 <= height*72 and all(font.getlength(line) <= width*72 for line in wrapped):
            break
    else:
        raise ValueError('Testo troppo lungo per la slide PowerPoint: '+value[:60])
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.text = value
    for paragraph in frame.paragraphs:
        paragraph.font.name = 'DejaVu Sans'
        paragraph.font.size = Pt(fitted)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor(*color)
        paragraph.line_spacing = 1.1
        paragraph.space_before = paragraph.space_after = Pt(0)
        if center:
            paragraph.alignment = PP_ALIGN.CENTER
        if link:
            for run in paragraph.runs:
                run.hyperlink.address = link
    return shape


def picture_contain(slide, path, x, y, width, height):
    with Image.open(path) as picture:
        ratio = min(width/picture.width, height/picture.height)
        fitted_width, fitted_height = picture.width*ratio, picture.height*ratio
    return slide.shapes.add_picture(str(path), Inches(x+(width-fitted_width)/2),
                                   Inches(y+(height-fitted_height)/2),
                                   width=Inches(fitted_width), height=Inches(fitted_height))


def brand(slide, x, y, width, white=False):
    image = logo_layer(560, white)
    stream = BytesIO()
    image.save(stream, format='PNG')
    stream.seek(0)
    slide.shapes.add_picture(stream, Inches(x), Inches(y), width=Inches(width))


def blank(deck, background):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(*background)
    return slide


def create_presentation(product, report, feed_path, destination, platform='instagram'):
    if platform not in ('instagram', 'facebook'):
        raise ValueError('Piattaforma social non supportata: '+platform)
    highlights = weekly_highlights(report)
    caption, tags = post_copy(product, report['headline'], report['cta'], report['source_claim'])
    start = datetime.fromisoformat(report['window_start']).strftime('%d/%m/%Y')
    end = datetime.fromisoformat(report['created_at']).strftime('%d/%m/%Y')
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    deck.core_properties.title = product['name']+' · Campagna IOMA'
    deck.core_properties.subject = 'Trend settimanali e anteprima del post '+platform.title()
    deck.core_properties.author = 'IOMA Paris · Beauty Agent'

    slide = blank(deck, MOCHA)
    brand(slide, .65, .45, 1.5, white=True)
    text(slide, 'CAMPAGNA BEAUTY', .65, 1.5, 6.1, .4, size=13, color=(221, 196, 164))
    text(slide, product['name'], .65, 2.05, 6.1, 2.1, size=32, minimum=24, bold=True, color=IVORY)
    rectangle(slide, .65, 4.35, .85, .035, GOLD)
    text(slide, report['headline'], .65, 4.7, 6.1, .8, size=24, color=IVORY)
    text(slide, 'Settimana '+start+' — '+end, .65, 6.55, 6.1, .45, size=13, color=IVORY)
    rectangle(slide, 7.4, 1.2, 5.25, 5.25, WHITE, rounded=True)
    picture_contain(slide, product['images'][0]['path'], 7.5, 1.3, 5.05, 5.05)
    slide.notes_slide.notes_text_frame.text = (
        'Titolo e claim dal catalogo locale. Packshot originale: '+product['images'][0]['source_url'])

    slide = blank(deck, IVORY)
    text(slide, 'I trend beauty della settimana', .6, .45, 10.2, .6, size=30, bold=True)
    text(slide, start+' — '+end, .6, 1.12, 10.2, .35, size=14, color=MUTED)
    brand(slide, 11.5, .55, 1.2)
    for index, article in enumerate(highlights):
        x = .6+index*4.13
        rectangle(slide, x, 1.85, 3.87, 4.65, WHITE, border=(222, 211, 197), rounded=True)
        rectangle(slide, x+.23, 2.12, .5, .035, GOLD)
        text(slide, article['topic'], x+.23, 2.35, 3.4, .9, size=20, minimum=16, bold=True)
        published = datetime.fromisoformat(article['published_at']).strftime('%d/%m/%Y')
        text(slide, 'Notizia del '+published, x+.23, 3.32, 3.4, .35, size=11, color=MUTED)
        text(slide, article['title'], x+.23, 3.85, 3.4, 1.8, size=17, minimum=12)
        text(slide, 'Fonte: '+article['publisher'], x+.23, 5.92, 3.4, .45,
             size=12, minimum=10, color=(116, 78, 36), link=article['url'])
    text(slide, 'Segnali editoriali dalle notizie; non misure di viralità sui social.',
         .6, 6.85, 12.1, .35, size=12, color=MUTED)
    slide.notes_slide.notes_text_frame.text = '\n\n'.join(
        f"{article['topic']}\n{article['title']}\n{article['publisher']} · {article['published_at']}\n{article['url']}"
        for article in highlights)

    slide = blank(deck, IVORY)
    text(slide, 'Anteprima del post · '+platform.title(), .6, .4, 10.6, .6, size=28, bold=True)
    brand(slide, 11.5, .5, 1.2)
    rectangle(slide, .6, 1.3, 12.1, 5.75, WHITE, border=(219, 213, 205), rounded=True)
    brand(slide, .83, 1.47, .75)
    text(slide, 'IOMA Paris', 1.82, 1.49, 5.2, .3, size=14, bold=True)
    text(slide, platform.title(), 10.8, 1.49, 1.62, .3, size=13, color=MUTED)
    rectangle(slide, .72, 1.95, 11.85, .012, (234, 229, 222))
    # Embed the full 4:5 export without cropping or stretching it.
    picture_contain(slide, feed_path, .8, 2.04, 3.92, 4.9)
    rectangle(slide, 5.03, 2.1, .012, 4.55, (234, 229, 222))
    text(slide, 'IOMA Paris', 5.42, 2.19, 6.75, .4, size=17, bold=True)
    text(slide, caption, 5.42, 2.9, 6.75, 2.65, size=19, minimum=14)
    text(slide, ' '.join(tags), 5.42, 5.78, 6.75, .75, size=15, color=(42, 93, 148))
    text(slide, 'Mi piace · Commenta · Condividi', 5.42, 6.65, 6.75, .3, size=12, color=MUTED)
    slide.notes_slide.notes_text_frame.text = (
        'Anteprima di layout, senza pubblicazione sul social.\n\nCaption:\n'+caption+
        '\n\nHashtag:\n'+' '.join(tags)+'\n\nImmagine: feed.png (1080×1350).')
    destination = Path(destination)
    deck.save(destination)
    return {'path': destination, 'slides': 3, 'platform': platform,
            'caption': caption, 'hashtags': tags, 'trends': highlights}
